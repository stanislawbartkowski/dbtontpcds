#!/usr/bin/env python3
"""Combine every benchmark result under result/ into one sheet.

Rows are the queries, followed by summary rows; columns are the targets,
grouped by data size, with the SQL engine description, data size and
per-query timeout in the header. Uses the newest result file for each
(data size, target) pair. Times are written as real Excel durations so
they can be sorted and summed; TIMEOUT and ERROR stay as text.

Usage:
    .venv/bin/python scripts/summarize_results.py [--result-dir result] [--out result/summary.xlsx]
"""
import argparse
import datetime
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FILE_NAME = re.compile(r"query_(?P<target>.+)_(?P<stamp>\d{8}_\d{6})\.xlsx$")
TIME_FORMAT = "[h]:mm:ss"
FAILED_FILL = PatternFill("solid", fgColor="F8D7DA")
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
SUMMARY_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")


def query_sort_key(name: str):
    m = re.match(r"query_(\d+)([a-z]?)", name)
    return (int(m.group(1)), m.group(2)) if m else (float("inf"), name)


def read_result(path: Path) -> dict:
    """Header info and {query: seconds | "TIMEOUT" | "ERROR"} from one result file.

    Older files have two info rows (no timeout row) instead of three, so the
    table header row is located by its "Query" cell."""
    raw = pd.read_excel(path, header=None)
    header_row = raw.index[raw[0] == "Query"][0]
    info = dict(str(v).split(": ", 1) for v in raw[0][:header_row] if ": " in str(v))
    table = pd.read_excel(path, skiprows=header_row)
    times = {}
    for query, value in zip(table["Query"], table["Execution Time"]):
        value = str(value)
        m = re.fullmatch(r"(\d+):(\d{2}):(\d{2})", value)
        times[query] = int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3]) if m else value
    return {
        "engine": info.get("SQL engine", ""),
        "size": info.get("Data size", ""),
        "timeout": info.get("Per-query timeout", "not recorded"),
        "times": times,
    }


def latest_results(result_dir: Path) -> list[dict]:
    newest = {}
    for path in result_dir.glob("*/query_*.xlsx"):
        m = FILE_NAME.search(path.name)
        if not m:
            continue
        key = (path.parent.name, m["target"])
        if key not in newest or m["stamp"] > newest[key][0]:
            newest[key] = (m["stamp"], path)
    columns = []
    for (size_dir, target), (_, path) in newest.items():
        columns.append({"size_dir": size_dir, "target": target, "file": path.name, **read_result(path)})
    columns.sort(key=lambda c: (int(c["size_dir"]) if c["size_dir"].isdigit() else 0, c["target"]))
    return columns


def duration(seconds: float) -> datetime.timedelta:
    return datetime.timedelta(seconds=round(seconds))


def write_summary(columns: list[dict], out_path: Path) -> None:
    queries = sorted({q for c in columns for q in c["times"]}, key=query_sort_key)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    labels = ["Data size", "Target", "SQL engine", "Per-query timeout", "Source file"]
    for r, label in enumerate(labels, start=1):
        ws.cell(r, 1, label).font = Font(bold=True)
    header_rows = len(labels)
    for i, c in enumerate(columns, start=2):
        for r, value in enumerate([c["size"], c["target"], c["engine"], c["timeout"], c["file"]], start=1):
            cell = ws.cell(r, i, value)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
        ws.cell(2, i).font = Font(bold=True)
    ws.cell(header_rows + 1, 1, "Query").font = Font(bold=True)

    row = header_rows + 2
    for q in queries:
        ws.cell(row, 1, q)
        for i, c in enumerate(columns, start=2):
            value = c["times"].get(q, "")
            cell = ws.cell(row, i)
            if isinstance(value, int):
                cell.value = duration(value)
                cell.number_format = TIME_FORMAT
            else:
                cell.value = value
                if value:
                    cell.fill = FAILED_FILL
            cell.alignment = Alignment(horizontal="center")
        row += 1

    # Averages over each engine's own completed queries aren't comparable when
    # engines fail different queries, so also average over the queries every
    # engine at that data size completed.
    common = {}
    for c in columns:
        done = {q for q, v in c["times"].items() if isinstance(v, int)}
        common[c["size_dir"]] = common.get(c["size_dir"], done) & done

    def completed(c):
        return [v for v in c["times"].values() if isinstance(v, int)]

    summary = [
        ("Completed", lambda c: len(completed(c)), None),
        ("Timeouts", lambda c: sum(v == "TIMEOUT" for v in c["times"].values()), None),
        ("Errors", lambda c: sum(v == "ERROR" for v in c["times"].values()), None),
        ("Total time (completed queries)", lambda c: duration(sum(completed(c))), TIME_FORMAT),
        ("Average time (completed queries)",
         lambda c: duration(sum(completed(c)) / len(completed(c))) if completed(c) else "", TIME_FORMAT),
        ("Queries completed by all targets at this size", lambda c: len(common[c["size_dir"]]), None),
        ("Average time (queries completed by all targets at this size)",
         lambda c: duration(sum(c["times"][q] for q in common[c["size_dir"]]) / len(common[c["size_dir"]]))
         if common[c["size_dir"]] else "", TIME_FORMAT),
    ]
    row += 1
    for label, fn, fmt in summary:
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 1).fill = SUMMARY_FILL
        for i, c in enumerate(columns, start=2):
            cell = ws.cell(row, i, fn(c))
            cell.fill = SUMMARY_FILL
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            if fmt:
                cell.number_format = fmt
        row += 1

    ws.column_dimensions["A"].width = 58
    for i in range(2, len(columns) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 24
    ws.row_dimensions[3].height = 75
    for r in range(1, row):
        for i in range(1, len(columns) + 2):
            ws.cell(r, i).border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.freeze_panes = ws.cell(header_rows + 2, 2)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--result-dir", default=str(PROJECT_ROOT / "result"))
    parser.add_argument("--out", default=str(PROJECT_ROOT / "result" / "summary.xlsx"))
    args = parser.parse_args()
    columns = latest_results(Path(args.result_dir))
    if not columns:
        raise SystemExit(f"no result files found under {args.result_dir}")
    write_summary(columns, Path(args.out))
    print(f"Wrote {len(columns)} result columns to {args.out}")
    for c in columns:
        print(f"  {c['size']:10} {c['target']:16} {c['file']}")


if __name__ == "__main__":
    main()
